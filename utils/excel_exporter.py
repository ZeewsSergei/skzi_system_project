import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Border, Side
from PyQt6.QtWidgets import QMessageBox
from logger import app_logger


class ExcelExporter:
    def __init__(self, db):
        self.db = db
        self.parent = None

    def set_parent(self, parent):
        self.parent = parent

    def fetch_dataframe(self, query, params=()):
        return pd.read_sql_query(query, self.db.conn, params=params)

    @staticmethod
    def _format_date_column(df, col_name):
        """Преобразует столбец с датами в формат dd.MM.yyyy, пустые значения остаются пустыми."""
        if col_name in df.columns:
            df[col_name] = pd.to_datetime(df[col_name], errors='coerce')
            df[col_name] = df[col_name].dt.strftime('%d.%m.%Y').fillna('')

    def _format_sheet(self, worksheet, df):
        """Автоподбор ширины, жирный заголовок, границы для заполненных ячеек."""
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Автоматическая ширина столбцов
        for col_idx, col_name in enumerate(df.columns):
            column_letter = get_column_letter(col_idx + 1)  # столбцы с A и т.д.
            max_len = max(
                df[col_name].astype(str).str.len().max(),
                len(col_name)
            ) + 2
            worksheet.column_dimensions[column_letter].width = max_len

        # Жирный заголовок и границы для всех ячеек с данными
        for row in worksheet.iter_rows(min_row=1, max_row=len(df)+1, max_col=len(df.columns)):
            for cell in row:
                if cell.row == 1:          # строка заголовка
                    cell.font = Font(bold=True)
                if cell.value is not None:
                    cell.border = thin_border

    def _write_and_format(self, df, file_path, sheet_name='Лист1'):
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            worksheet = writer.sheets[sheet_name]
            self._format_sheet(worksheet, df)

    # ---------- акт установки ----------
    def export_install_act(self, start_date, end_date, file_path):
        query = """
        SELECT
            e.fio AS 'Фамилия, имя, отчество',
            a.cabinet_number AS 'Номер помещения и рабочего места',
            a.arm_serial AS 'Заводской номер АРМ',
            sn.name AS 'Наименование и версия СКЗИ',
            r.install_date AS 'Дата установки'
        FROM skzi_registry r
        JOIN employees e ON r.employee_id = e.id
        JOIN arm a ON r.arm_id = a.id
        LEFT JOIN skzi_names sn ON r.skzi_name_id = sn.id
        WHERE r.status = 'ACTIVE'
        """
        params = []
        if start_date and end_date:
            query += " AND r.install_date BETWEEN ? AND ?"
            params = [start_date, end_date]
        query += " ORDER BY r.install_date"

        df = self.fetch_dataframe(query, params)
        if df.empty:
            if self.parent:
                QMessageBox.warning(self.parent, "Нет данных",
                                    "За выбранный период нет активных записей СКЗИ.\nФайл не будет создан.")
            return

        self._format_date_column(df, 'Дата установки')
        df.insert(0, '№ п/п', range(1, len(df)+1))
        self._write_and_format(df, file_path)
        app_logger.info(f"Экспортирован акт установки: {file_path}, записей: {len(df)}")

    # ---------- акт уничтожения ----------
    def export_destruction_act(self, start_date, end_date, file_path):
        iso_withdrawal = (
            "CASE WHEN length(r.withdrawal_date) = 10 "
            "THEN substr(r.withdrawal_date, 7, 4) || '-' || substr(r.withdrawal_date, 4, 2) || '-' || substr(r.withdrawal_date, 1, 2) "
            "ELSE r.withdrawal_date END"
        )
        iso_vip_withdrawal = (
            "CASE WHEN length(v.withdrawal_date) = 10 "
            "THEN substr(v.withdrawal_date, 7, 4) || '-' || substr(v.withdrawal_date, 4, 2) || '-' || substr(v.withdrawal_date, 1, 2) "
            "ELSE v.withdrawal_date END"
        )

        skzi_sql = f"""
        SELECT
            r.media_number,
            r.skzi_number,
            e.fio,
            1,
            r.skzi_instance_number,
            1,
            '' as note,
            {iso_withdrawal} AS withdrawal_date
        FROM skzi_registry r
        JOIN employees e ON r.employee_id = e.id
        WHERE r.status = 'DESTROYED'
        """
        vipnet_sql = f"""
        SELECT
            '',
            v.skzi_account,
            e.fio,
            1,
            '',
            1,
            '' as note,
            {iso_vip_withdrawal} AS withdrawal_date
        FROM vipnet_installations v
        JOIN employees e ON v.employee_id = e.id
        WHERE v.status = 'DESTROYED'
        """

        params = []
        if start_date:
            date_cond = f" AND withdrawal_date >= ?"
            skzi_sql += date_cond
            vipnet_sql += date_cond
            params.append(start_date)
        if end_date:
            date_cond = f" AND withdrawal_date <= ?"
            skzi_sql += date_cond
            vipnet_sql += date_cond
            params.append(end_date)
        union_params = params * 2

        union_query = f"""
        SELECT * FROM (
            {skzi_sql}
            UNION ALL
            {vipnet_sql}
        ) ORDER BY withdrawal_date
        """

        df = self.fetch_dataframe(union_query, union_params)
        df.columns = [
            'Учетный номер ключевого носителя',
            'Номер криптографического ключа, наименование документа',
            'Владелец ключа (документа)',
            'Количество ключевых носителей',
            'Номера экземпляров',
            'Всего ключей',
            'Примечание',
            'Дата уничтожения'
        ]
        self._format_date_column(df, 'Дата уничтожения')

        if df.empty:
            if self.parent:
                QMessageBox.warning(self.parent, "Нет данных",
                                    "За выбранный период нет уничтоженных записей (СКЗИ и ViPNet).")
            return

        df.insert(0, '№ п/п', range(1, len(df)+1))
        self._write_and_format(df, file_path)
        app_logger.info(f"Экспортирован акт уничтожения: {file_path}, записей: {len(df)}")

    # ---------- ведомость обучения ----------
    def export_training_sheet(self, start_date, end_date, file_path):
        iso_train = (
            "CASE WHEN length(t.training_date) = 10 "
            "THEN substr(t.training_date, 7, 4) || '-' || substr(t.training_date, 4, 2) || '-' || substr(t.training_date, 1, 2) "
            "ELSE t.training_date END"
        )
        query = f"""
        SELECT
            e.fio AS 'Фамилия, имя, отчество',
            t.position AS 'Должность',
            t.department AS 'Отдел',
            t.result AS 'Отметка о результатах проверки знаний'
        FROM training_logs t
        JOIN employees e ON t.employee_id = e.id
        """
        params = []
        if start_date and end_date:
            query += f" WHERE {iso_train} BETWEEN ? AND ?"
            params = [start_date, end_date]
        query += f" ORDER BY {iso_train}"

        df = self.fetch_dataframe(query, params)
        if df.empty:
            if self.parent:
                QMessageBox.warning(self.parent, "Нет данных",
                                    "За выбранный период нет записей об обучении.\nФайл не будет создан.")
            return

        df.insert(0, '№ п/п', range(1, len(df)+1))
        self._write_and_format(df, file_path)
        app_logger.info(f"Экспортирована ведомость обучения: {file_path}, записей: {len(df)}")

    # ---------- полный учёт ----------
    def export_full_data(self, file_path):
        iso_withdrawal_reg = (
            "CASE WHEN length(r.withdrawal_date) = 10 "
            "THEN substr(r.withdrawal_date, 7, 4) || '-' || substr(r.withdrawal_date, 4, 2) || '-' || substr(r.withdrawal_date, 1, 2) "
            "ELSE r.withdrawal_date END"
        )
        iso_train = (
            "CASE WHEN length(t.training_date) = 10 "
            "THEN substr(t.training_date, 7, 4) || '-' || substr(t.training_date, 4, 2) || '-' || substr(t.training_date, 1, 2) "
            "ELSE t.training_date END"
        )
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            # Реестр
            df_reg = self.fetch_dataframe(f"""
                SELECT r.id,
                       e.fio as 'Сотрудник',
                       sn.name as 'СКЗИ',
                       r.skzi_number as 'Заводской №',
                       r.skzi_instance_number as '№ экземпляра',
                       a.cabinet_number as 'Кабинет',
                       r.install_date as 'Дата установки',
                       {iso_withdrawal_reg} as 'Дата изъятия',
                       r.status as 'Статус'
                FROM skzi_registry r
                JOIN employees e ON r.employee_id = e.id
                JOIN arm a ON r.arm_id = a.id
                LEFT JOIN skzi_names sn ON r.skzi_name_id = sn.id
            """)
            self._format_date_column(df_reg, 'Дата установки')
            self._format_date_column(df_reg, 'Дата изъятия')
            df_reg.to_excel(writer, sheet_name='Реестр', index=False)
            self._format_sheet(writer.sheets['Реестр'], df_reg)

            # Сотрудники
            df_emp = self.fetch_dataframe("""
                SELECT e.fio as 'ФИО',
                       e.position as 'Должность',
                       s.name as 'Сектор',
                       d.name as 'Отдел',
                       e.knowledge_check as 'Проверка знаний'
                FROM employees e
                LEFT JOIN sectors s ON e.sector_id = s.id
                LEFT JOIN departments d ON e.department_id = d.id
            """)
            df_emp.to_excel(writer, sheet_name='Сотрудники', index=False)
            self._format_sheet(writer.sheets['Сотрудники'], df_emp)

            # АРМ
            df_arm = self.fetch_dataframe("""
                SELECT a.id,
                       a.arm_name as 'Имя АРМ',
                       a.arm_serial as 'Серийный №',
                       at.name as 'Тип АРМ',
                       a.cabinet_number as 'Кабинет',
                       addr.name as 'Адрес установки'
                FROM arm a
                LEFT JOIN arm_types at ON a.arm_type_id = at.id
                LEFT JOIN addresses addr ON a.install_address_id = addr.id
            """)
            df_arm.to_excel(writer, sheet_name='АРМ', index=False)
            self._format_sheet(writer.sheets['АРМ'], df_arm)

            # Обучение
            df_train = self.fetch_dataframe(f"""
                SELECT {iso_train} as 'Дата',
                       e.fio as 'Сотрудник',
                       t.position as 'Должность',
                       t.department as 'Отдел',
                       t.result as 'Результат'
                FROM training_logs t
                JOIN employees e ON t.employee_id = e.id
            """)
            self._format_date_column(df_train, 'Дата')
            df_train.to_excel(writer, sheet_name='Обучение', index=False)
            self._format_sheet(writer.sheets['Обучение'], df_train)

        app_logger.info(f"Экспортированы все данные: {file_path}")

    # ---------- журнал ЭП ----------
    def export_ep_journal(self, file_path):
        iso_withdrawal = (
            "CASE WHEN length(r.withdrawal_date) = 10 "
            "THEN substr(r.withdrawal_date, 7, 4) || '-' || substr(r.withdrawal_date, 4, 2) || '-' || substr(r.withdrawal_date, 1, 2) "
            "ELSE r.withdrawal_date END"
        )
        query = f"""
        SELECT
            mt.name AS 'Тип носителя',
            r.media_number AS 'Номер носителя ЭП',
            r.cert_number AS '№ сертификата ЭП',
            rf.name AS 'От кого получен',
            r.receive_letter_num AS '№ письма',
            r.receive_date AS 'Дата получения письма',
            e.fio AS 'ФИО',
            d.name AS 'Отдел',
            r.install_date AS 'Дата установки',
            r.installer_fio AS 'Кто установил',
            a.cabinet_number AS 'Кабинет',
            a.arm_serial AS 'Серийный № АРМ',
            {iso_withdrawal} AS 'Дата уничтожения',
            r.withdrawer_fio AS 'ФИО изъявшего',
            r.destruction_act_num AS '№ акта уничтожения',
            r.expiry_date AS 'Срок до'
        FROM skzi_registry r
        LEFT JOIN employees e ON r.employee_id = e.id
        LEFT JOIN departments d ON e.department_id = d.id
        LEFT JOIN arm a ON r.arm_id = a.id
        LEFT JOIN media_types mt ON r.media_type_id = mt.id
        LEFT JOIN received_from rf ON r.received_from_id = rf.id
        WHERE r.media_number IS NOT NULL AND r.media_number != ''
        ORDER BY r.install_date DESC
        """
        df = self.fetch_dataframe(query)
        if df.empty:
            if self.parent:
                QMessageBox.warning(self.parent, "Нет данных", "Нет записей для экспорта журнала ЭП.")
            return
        self._format_date_column(df, 'Дата получения письма')
        self._format_date_column(df, 'Дата установки')
        self._format_date_column(df, 'Дата уничтожения')
        self._format_date_column(df, 'Срок до')
        df.insert(0, '№ п/п', range(1, len(df)+1))
        self._write_and_format(df, file_path)
        app_logger.info(f"Экспортирован журнал ЭП: {file_path}, записей: {len(df)}")

    # ---------- журнал СКЗИ ----------
    def export_skzi_journal(self, file_path):
        iso_withdrawal = (
            "CASE WHEN length(r.withdrawal_date) = 10 "
            "THEN substr(r.withdrawal_date, 7, 4) || '-' || substr(r.withdrawal_date, 4, 2) || '-' || substr(r.withdrawal_date, 1, 2) "
            "ELSE r.withdrawal_date END"
        )
        query = f"""
        SELECT
            sn.name AS 'Наименование СКЗИ',
            r.skzi_number AS 'Серийный (заводской) № СКЗИ',
            r.skzi_instance_number AS '№ экземпляра',
            rf.name AS 'От кого получен',
            r.receive_letter_num AS '№ письма',
            r.receive_date AS 'Дата получения письма',
            e.fio AS 'ФИО',
            r.install_date AS 'Дата установки',
            r.installer_fio AS 'Кто установил',
            a.cabinet_number AS 'Кабинет',
            a.arm_serial AS 'Серийный № АРМ',
            {iso_withdrawal} AS 'Дата уничтожения',
            r.withdrawer_fio AS 'ФИО изъявшего',
            r.destruction_act_num AS '№ акта уничтожения',
            r.expiry_date AS 'Срок до'
        FROM skzi_registry r
        LEFT JOIN employees e ON r.employee_id = e.id
        LEFT JOIN arm a ON r.arm_id = a.id
        LEFT JOIN skzi_names sn ON r.skzi_name_id = sn.id
        LEFT JOIN received_from rf ON r.received_from_id = rf.id
        WHERE r.skzi_number IS NOT NULL AND r.skzi_number != ''
        ORDER BY r.install_date DESC
        """
        df = self.fetch_dataframe(query)
        if df.empty:
            if self.parent:
                QMessageBox.warning(self.parent, "Нет данных", "Нет записей для экспорта журнала СКЗИ.")
            return
        self._format_date_column(df, 'Дата получения письма')
        self._format_date_column(df, 'Дата установки')
        self._format_date_column(df, 'Дата уничтожения')
        self._format_date_column(df, 'Срок до')
        df.insert(0, '№ п/п', range(1, len(df)+1))
        self._write_and_format(df, file_path)
        app_logger.info(f"Экспортирован журнал СКЗИ: {file_path}, записей: {len(df)}")

    # ---------- журнал СЗИ от НСД ----------
    def export_szi_nsd_journal(self, file_path):
        query = """
        SELECT
            e.fio AS 'ФИО',
            d.name AS 'Отдел',
            sz.name AS 'Наименование СЗИ от НСД',
            a.arm_name AS 'Имя АРМ',
            a.arm_serial AS 'Серийный номер АРМ',
            a.cabinet_number AS 'Кабинет',
            addr.name AS 'Адрес установки',
            s.install_date AS 'Дата установки'
        FROM szi_nsd_installations s
        LEFT JOIN employees e ON s.employee_id = e.id
        LEFT JOIN departments d ON e.department_id = d.id
        LEFT JOIN arm a ON s.arm_id = a.id
        LEFT JOIN szi_nsd_names sz ON s.szi_nsd_id = sz.id
        LEFT JOIN addresses addr ON a.install_address_id = addr.id
        WHERE s.status = 'ACTIVE'
        ORDER BY s.install_date DESC
        """
        df = self.fetch_dataframe(query)
        if df.empty:
            if self.parent:
                QMessageBox.warning(self.parent, "Нет данных",
                                    "Нет активных записей СЗИ от НСД для экспорта.\nФайл не будет создан.")
            return
        self._format_date_column(df, 'Дата установки')
        df.insert(0, '№ п/п', range(1, len(df)+1))
        self._write_and_format(df, file_path)
        app_logger.info(f"Экспортирован журнал СЗИ от НСД: {file_path}, записей: {len(df)}")

    # ---------- журнал ViPNet Client ----------
    def export_vipnet_journal(self, file_path):
        # Формулы для преобразования дат из dd.MM.yyyy в yyyy-MM-dd
        iso_install = (
            "CASE WHEN length(v.install_date) = 10 "
            "THEN substr(v.install_date, 7, 4) || '-' || substr(v.install_date, 4, 2) || '-' || substr(v.install_date, 1, 2) "
            "ELSE v.install_date END"
        )
        iso_withdrawal = (
            "CASE WHEN length(v.withdrawal_date) = 10 "
            "THEN substr(v.withdrawal_date, 7, 4) || '-' || substr(v.withdrawal_date, 4, 2) || '-' || substr(v.withdrawal_date, 1, 2) "
            "ELSE v.withdrawal_date END"
        )
        query = f"""
        SELECT
            sn.name AS skzi_name,
            v.skzi_account,
            rf.name AS received_from,
            v.receive_letter_num AS letter_num,
            e.fio,
            d.name AS department,
            {iso_install} AS install_date1,
            a.arm_name,
            a.arm_serial,
            at.name AS arm_type,
            a.cabinet_number AS cabinet,
            addr.name AS install_address,
            v.installer_fio,
            {iso_withdrawal} AS withdrawal_date,
            v.withdrawer_fio,
            v.destruction_act_num AS act_num
        FROM vipnet_installations v
        LEFT JOIN employees e ON v.employee_id = e.id
        LEFT JOIN departments d ON e.department_id = d.id
        LEFT JOIN arm a ON v.arm_id = a.id
        LEFT JOIN arm_types at ON a.arm_type_id = at.id
        LEFT JOIN skzi_names sn ON v.skzi_name_id = sn.id
        LEFT JOIN received_from rf ON v.received_from_id = rf.id
        LEFT JOIN addresses addr ON a.install_address_id = addr.id
        ORDER BY v.install_date DESC
        """
        df = self.fetch_dataframe(query)
        if df.empty:
            if self.parent:
                QMessageBox.warning(self.parent, "Нет данных", "Нет записей ViPNet Client для экспорта.")
            return
        df.columns = [
            'Наименование СКЗИ',
            'Узел ViPNet Client',
            'От кого получен',
            '№ письма',
            'ФИО',
            'Отдел',
            'Дата установки',
            'Имя АРМ',
            'Серийный № АРМ',
            'Тип АРМ',
            'Кабинет',
            'Адрес установки',
            'Кто установил',
            'Дата уничтожения',
            'ФИО изъявшего',
            '№ акта уничтожения'
        ]
        self._format_date_column(df, 'Дата установки')
        self._format_date_column(df, 'Дата уничтожения')
        df.insert(0, '№ п/п', range(1, len(df)+1))
        self._write_and_format(df, file_path)
        app_logger.info(f"Экспортирован журнал ViPNet Client: {file_path}, записей: {len(df)}")

    def export_vipnet_install_act(self, start_date, end_date, file_path):
        """
        Экспорт акта установки ViPNet Client.
        Столбцы: № п/п, Фамилия, имя, отчество, Кабинет, Серийный № АРМ,
                 СКЗИ, Узел ViPNet Client, Дата установки.
        """
        iso_install = (
            "CASE WHEN length(v.install_date) = 10 "
            "THEN substr(v.install_date, 7, 4) || '-' || substr(v.install_date, 4, 2) || '-' || substr(v.install_date, 1, 2) "
            "ELSE v.install_date END"
        )
        query = f"""
        SELECT
            e.fio AS 'Фамилия, имя, отчество',
            a.cabinet_number AS 'Кабинет',
            a.arm_serial AS 'Серийный № АРМ',
            sn.name AS 'СКЗИ',
            v.skzi_account AS 'Узел ViPNet Client',
            {iso_install} AS 'Дата установки'
        FROM vipnet_installations v
        JOIN employees e ON v.employee_id = e.id
        JOIN arm a ON v.arm_id = a.id
        LEFT JOIN skzi_names sn ON v.skzi_name_id = sn.id
        WHERE v.status = 'ACTIVE'
        """
        params = []
        if start_date and end_date:
            query += f" AND {iso_install} BETWEEN ? AND ?"
            params = [start_date, end_date]
        query += " ORDER BY v.install_date"

        df = self.fetch_dataframe(query, params)
        if df.empty:
            if self.parent:
                from PyQt6.QtWidgets import QMessageBox
                QMessageBox.warning(self.parent, "Нет данных", "За выбранный период нет активных ViPNet Client.")
            return

        self._format_date_column(df, 'Дата установки')
        df.insert(0, '№ п/п', range(1, len(df) + 1))
        # Переупорядочим столбцы при необходимости
        df = df[['№ п/п', 'Фамилия, имя, отчество', 'Кабинет', 'Серийный № АРМ', 'СКЗИ', 'Узел ViPNet Client',
                 'Дата установки']]
        self._write_and_format(df, file_path)
        from logger import app_logger
        app_logger.info(f"Экспортирован акт установки ViPNet Client: {file_path}, записей: {len(df)}")